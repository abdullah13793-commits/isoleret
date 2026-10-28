"""Audit report (PDF) and workbook (XLSX).

PDF: summary, method and validation, the runs with status and proposal, fire cross-check, material list, charts.
python -m isoleret.report  ->  out/Isoleringsrapport_blok_c.pdf, out/Isoleringsaudit_blok_c.xlsx
"""
import csv
import io
import json
import math
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import ds452, physics
from . import fonts

ROOT = Path(__file__).resolve().parents[2]
STATUS = {"OK": "Opfylder DS 452", "UNDER": "For tynd isolering", "MANGLER": "Uisoleret", "KONDENS": "Kondensrisiko", "DS439": "Frostsikring (DS 439)",
          "INGEN_KRAV": "Intet krav"}
COL = {"OK": "#16a34a", "UNDER": "#d97706", "MANGLER": "#dc2626", "KONDENS": "#7c3aed", "DS439": "#6b7280", "INGEN_KRAV": "#9ca3af"}
SYS = {"heating": "Varme", "dhw": "Varmt brugsvand", "cold_water": "Koldt vand"}
INK, MUTED = colors.HexColor("#111827"), colors.HexColor("#6B7280")


def _fonts():
    pdfmetrics.registerFont(TTFont("Seg", fonts.path()))
    pdfmetrics.registerFont(TTFont("SegB", fonts.path(True)))


def _chart_savings(runs):
    rows = sorted([r for r in runs if r["proposal"] and r["proposal"]["kwh"] > 0], key=lambda r: -r["proposal"]["kwh"])[:12]
    fig, ax = plt.subplots(figsize=(7.2, 3.2), dpi=200)
    ax.barh([r["id"] for r in rows][::-1], [r["proposal"]["kwh"] for r in rows][::-1], color="#2f63c4", height=0.6)
    for i, r in enumerate(rows[::-1]):
        ax.text(r["proposal"]["kwh"] + 20, i, (f"{r['proposal']['kwh']:,} kWh · ".replace(",", ".") + (f"{r['proposal']['payback_years']:.1f}".replace(".", ",") if r["proposal"]["payback_years"] else "–") + " år"), va="center", fontsize=7, color="#374151")
    ax.set_xlabel("Besparelse, kWh pr. år", fontsize=8)
    ax.tick_params(labelsize=7)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_xlim(0, max(r["proposal"]["kwh"] for r in rows) * 1.35)
    fig.tight_layout()
    b = io.BytesIO()
    fig.savefig(b, format="png")
    plt.close(fig)
    b.seek(0)
    return b


def _chart_validation(val):
    fig, axs = plt.subplots(1, 2, figsize=(7.2, 3.0), dpi=200)
    v1 = val["V1_DS452_annexA"]["data"]
    for tab, c in (("A.1", "#2f63c4"), ("A.2", "#d97706")):
        d = [x for x in v1 if x["table"] == tab]
        axs[0].scatter([x["published"] for x in d], [x["model"] for x in d], s=10, color=c, label=f"tabel {tab}")
    axs[0].plot([0, 180], [0, 180], color="#9ca3af", lw=0.8)
    axs[0].set_title("DS 452 anneks A: isoleringstykkelse, mm", fontsize=8)
    axs[0].set_xlabel("standarden", fontsize=7)
    axs[0].set_ylabel("Isoleret", fontsize=7)
    axs[0].legend(fontsize=6, frameon=False)
    v2 = val["V2_Energihaandbogen"]["data"]
    axs[1].scatter([x["published"] for x in v2], [x["model"] for x in v2], s=10, color="#2f63c4", label="uden 50 %-regel")
    bare = [x for x in v2 if x["model_with_halving"] is not None and x["from_mm"] == 0]
    axs[1].scatter([x["published"] for x in bare], [x["model_with_halving"] for x in bare], s=10, color="#dc2626", marker="x", label="med 50 %-regel (kælder)")
    axs[1].plot([0, 1450], [0, 1450], color="#9ca3af", lw=0.8)
    axs[1].set_title("Energihåndbogen: besparelse, kWh/m/år", fontsize=8)
    axs[1].set_xlabel("Energihåndbogen", fontsize=7)
    axs[1].legend(fontsize=6, frameon=False)
    for a in axs:
        a.tick_params(labelsize=7)
        for s in ("top", "right"):
            a.spines[s].set_visible(False)
    fig.tight_layout()
    b = io.BytesIO()
    fig.savefig(b, format="png")
    plt.close(fig)
    b.seek(0)
    return b


def pdf(audit, val, out):
    _fonts()
    st = {k: ParagraphStyle(k, fontName=f, fontSize=s, leading=s * 1.35, textColor=INK, spaceAfter=a)
          for k, f, s, a in [("b", "Seg", 8.5, 3), ("h", "SegB", 12, 5), ("t", "SegB", 20, 4), ("c", "Seg", 7, 0), ("cb", "SegB", 7, 0), ("m", "Seg", 8, 2)]}
    P = lambda t, s="b": Paragraph(t, st[s])
    s = audit["summary"]
    story = [P("Isoleringsaudit, Blok C", "t"), P("Varme-, brugsvands- og koldtvandsanlæg · DS 452 · EN ISO 12241 · fiktivt demonstrationsprojekt", "m"), Spacer(1, 6)]
    kpi = [["Strækninger / meter rør", f"{s['runs']} / {s['metres']:.0f} m"], ["Opfylder DS 452", str(s["status"].get("OK", 0))],
           ["For tynd / uisoleret / kondensrisiko", f"{s['status'].get('UNDER', 0)} / {s['status'].get('MANGLER', 0)} / {s['status'].get('KONDENS', 0)}"],
           ["Varmetab i dag (rør og komponenter)", f"{s['kwh_now']:,} kWh/år".replace(",", ".")],
           ["Besparelse ved forslaget", f"{s['kwh_saved']:,} kWh/år · {s['dkk_per_year']:,} kr./år · {s['co2_kg']:,} kg CO₂/år".replace(",", ".")],
           ["Overslag, materialer og montage", f"{s['cost_dkk']:,} kr. · simpel tilbagebetaling {str(s['payback_years']).replace('.', ',')} år".replace(",", ".")],
           ["Brandgennemføringer berørt", f"{s['penetrations']}, heraf {s['fire_conflicts']} hvor efterisolering ville forlade den testede opbygning"],
           ["Uisolerede komponenter", f"{s['components_uninsulated']} (ventiler, pumper, målere)"]]
    t = Table([[P(a), P(f"<b>{b}</b>")] for a, b in kpi], colWidths=[75 * mm, 100 * mm])
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#E5E7EB"))]))
    story += [t, Spacer(1, 8), Image(_chart_savings(audit["runs"]), width=175 * mm, height=78 * mm),
              P("Største besparelser pr. strækning med simpel tilbagebetalingstid. Priser og CO₂-faktorer er eksempler (fjernvarme 0,75 kr./kWh, "
                "0,066 kg CO₂/kWh) og rettes til projektets tal.", "c"),
              Spacer(1, 6), P("Metode", "h"),
              P("Isoleringsklasse fra DS 452 tabel 6.4 ud fra anlæg, installationsdel, drift og omgivelsestemperatur (≤ 5 °C / > 5 °C). "
                "Varmetab efter EN ISO 12241 for cylindriske lag med λ ved lagets middeltemperatur og overgangskoefficient fra konvektion og stråling. "
                "Bæringer, der gennembryder isoleringen, giver +15 % (bygninger) / +25 % (ude). Komponenter: halv tykkelse (tabel 6.2); "
                "en ventil regnes som 0,2 m rør, en pumpe som 1 m (Energihåndbogen). Koldt vand kontrolleres for dugpunkt (Magnus)."),
              P("Validering", "h"),
              P(f"DS 452 anneks A, tabel A.1/A.2 (72 værdier): genskabt med RMS {val['V1_DS452_annexA']['rms_mm']} mm, max {val['V1_DS452_annexA']['max_abs_mm']} mm "
                f"(design-λ {val['V1_DS452_annexA']['fitted_lambda']} W/mK tilpasset, da standarden ikke angiver den). "
                f"Energihåndbogen tabel 11/12/15/16 (96 værdier, λ 0,038 som angivet): uisolerede rækker inden for {val['V2_Energihaandbogen']['bare_max_rel_pct']} %, "
                f"øvrige gennemsnitlig {val['V2_Energihaandbogen']['insulated_mean_abs_kwh']} kWh. Bemærk: tabellerne for kælder angiver, at halvdelen af tabet udnyttes, "
                f"men tallene genskabes kun uden den regel (med reglen ville alle kælderværdier være {abs(val['V2_Energihaandbogen']['halving_finding']['mean_rel_err_with_halving_pct']):.0f} % lavere)."),
              Image(_chart_validation(val), width=175 * mm, height=73 * mm), PageBreak()]
    head = ["ID", "Strækning", "Rum", "Kl.", "Ø", "m", "Isolering nu", "U / grænse", "Status", "Forslag", "kWh/år", "Tilbage-bet."]
    rows = [[P(h, "cb") for h in head]]
    for r in audit["runs"]:
        ins = r["insulation"]
        p = r["proposal"]
        rows.append([P(r["id"], "c"), P(r["name"], "c"), P(f"{r['space']} ({r['ambient']:g} °C)", "c"), P(str(r["class"]), "c"), P(f"{r['od_mm']:g}", "c"),
                     P(f"{r['length']:.1f}", "c"), P(f"{(ins or {}).get('product') or '–'} {round(ins['thickness'] * 1000) if ins else ''}", "c"),
                     P(f"{r['U_now']:.3f} / {r['U_limit']:.3f}" if r["U_limit"] else "–", "c"),
                     P(f"<font color='{COL[r['status']]}'><b>{STATUS[r['status']]}</b></font>", "c"),
                     P(f"{p['product']} {p['thickness_mm']} mm" if p else "–", "c"), P(f"{p['kwh']:,}".replace(",", ".") if p else "–", "c"),
                     P(f"{p['payback_years']:.1f} år".replace(".", ",") if p and p["payback_years"] else "–", "c")])
    tb = Table(rows, colWidths=[13 * mm, 52 * mm, 30 * mm, 8 * mm, 10 * mm, 10 * mm, 34 * mm, 24 * mm, 26 * mm, 34 * mm, 16 * mm, 17 * mm], repeatRows=1)
    tb.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6"))]))
    story += [P("Strækninger", "h"), tb, PageBreak(), P("Brandkrydstjek: isolering i brandgennemføringer", "h"),
              P("En gennemføringstætning er kun dokumenteret med den isolering, den er prøvet med. Hvor efterisolering ville ændre isoleringen i "
                "tætningen, anbefales det at beholde den testede isolering over den testede længde; DS 452 tabel 6.2 kræver kun halv tykkelse i gennemføringer.")]
    frows = [[P(h, "cb") for h in ["Strækning", "Konstruktion", "I dag", "Efter forslag", "Anbefaling"]]]
    for r in audit["runs"]:
        for x in r["fire"]:
            ref = lambda v: f"{v['status']}" + (f"<br/>{v['ref']['eta']} {v['ref']['section']} s. {v['ref']['page']}" if v.get("ref") else "")
            frows.append([P(f"{r['id']} · {r['name']}", "c"), P(f"{x['host']} ({x['rating']})", "c"), P(ref(x["now"]), "c"), P(ref(x["new"]), "c"),
                          P(x["recommendation"], "c")])
    ft = Table(frows, colWidths=[55 * mm, 38 * mm, 36 * mm, 36 * mm, 105 * mm], repeatRows=1)
    ft.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6"))]))
    story += [ft, PageBreak(), P("Materialeliste (forslag)", "h")]
    brows = [[P(h, "cb") for h in ["Produkt", "Rør Ø mm", "Tykkelse mm", "Kappe", "Meter", "Bestilles"]]]
    for b in audit["bom"]:
        brows.append([P(b["product"], "c"), P(f"{b['pipe_od_mm']:g}", "c"), P(str(b["thickness_mm"]), "c"), P(b["cladding"], "c"),
                      P(f"{b.get('metres', '')}", "c"), P(f"{b['order_m']} m" if "order_m" in b else f"{b['pieces']} stk.", "c")])
    bt = Table(brows, colWidths=[80 * mm, 22 * mm, 22 * mm, 25 * mm, 20 * mm, 22 * mm], repeatRows=1)
    bt.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB"))]))
    story += [bt, Spacer(1, 6), P("Varenumre efter leverandørens sortiment. +5 % til spild. Demonstrationsprojektet er fiktivt; "
                                  "λ-værdier er typiske deklarerede værdier for materialeklassen og erstattes af det valgte produkts data.", "c")]
    story += [PageBreak(), P("Bilag: beregningstrin pr. strækning (i dag)", "h"),
              P("q = (θm − θa) / (Σ ln(Do/Di)/(2πλ) + 1/(π·De·hse)), λ ved lagets middeltemperatur; hse = hcv + hr; U = q·f/(θm − θa), f = tillæg for bæringer. "
                "Alle værdier kan regnes efter i hånden; appen viser de samme trin under 'Vis beregning'.", "c")]
    crow = [[P(h, "cb") for h in ["ID", "D m", "θm/θa °C", "Lag (mm, θmiddel, λ, R)", "Ts °C", "hse W/m²K", "Rse mK/W", "q W/m", "f", "U W/mK", "Grænse"]]]
    for r in audit["runs"]:
        ins = r["insulation"]
        lay = [(ins["thickness"], ins["material"])] if ins else []
        res = physics.pipe(r["od_mm"] / 1000, lay, r["tm"], r["ambient"], (ins or {}).get("cladding", "laerred"), bare={"Kobber": "copper"}.get(r["material"], "steel"))
        if lay:
            tmid = r["tm"] - res["q"] * res["R"][0] / 2
            lt = f"{ins['thickness'] * 1000:.0f} mm, ≈{tmid:.1f} °C, λ {res['lam'][0]:.4f}, R {res['R'][0]:.4f}"
            rse = 1 / (math.pi * res["De"] * res["h_se"])
        else:
            lt, rse = "uisoleret", None
        f = r["support_factor"] if lay else 1.0
        crow.append([P(r["id"], "c"), P(f"{r['od_mm'] / 1000:.4f}", "c"), P(f"{r['tm']:g}/{r['ambient']:g}", "c"), P(lt, "c"), P(f"{res['ts']:.2f}", "c"),
                     P(f"{res['h_se']:.3f}", "c"), P(f"{rse:.4f}" if rse else "–", "c"), P(f"{res['q']:.2f}", "c"), P(f"{f:g}", "c"),
                     P(f"{res['q'] * f / (r['tm'] - r['ambient']):.4f}", "c"), P(f"{r['U_limit']:.4f}" if r["U_limit"] else "–", "c")])
    ct = Table(crow, colWidths=[13 * mm, 15 * mm, 18 * mm, 72 * mm, 14 * mm, 18 * mm, 18 * mm, 15 * mm, 9 * mm, 18 * mm, 18 * mm], repeatRows=1)
    ct.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6"))]))
    story += [ct, Spacer(1, 4), P("θmiddel = (θm + Ts)/2, den temperatur λ er bestemt ved (løst iterativt). Kontrol: q = (θm − θa)/(R + Rse).", "c")]
    SimpleDocTemplate(str(out), pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                      title="Isoleringsaudit, Blok C", author="").build(story)


def xlsx(audit, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Strækninger"
    head = ["ID", "Strækning", "System", "Del", "Rum", "Omg. °C", "Medie °C", "Timer", "Ø mm", "Længde m", "Klasse", "Isolering", "mm", "U W/mK",
            "Grænse W/mK", "Status", "Forslag", "Forslag mm", "kWh/år", "kr./år", "Pris kr.", "Tilbagebetaling år", "kg CO2/år"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in audit["runs"]:
        ins, p = r["insulation"] or {}, r["proposal"] or {}
        ws.append([r["id"], r["name"], SYS.get(r["system"], r["system"]), ds452.LABEL.get(r["part"], r["part"]), r["space"], r["ambient"], r["tm"], r["hours"],
                   r["od_mm"], r["length"], str(r["class"]), ins.get("product"), round(ins["thickness"] * 1000) if ins.get("thickness") else None,
                   r["U_now"], r["U_limit"], STATUS[r["status"]], p.get("product"), p.get("thickness_mm"), p.get("kwh"), p.get("dkk_per_year"),
                   p.get("cost_dkk"), p.get("payback_years"), p.get("co2_kg")])
        ws.cell(ws.max_row, 16).fill = PatternFill("solid", fgColor=COL[r["status"]].lstrip("#"))
        ws.cell(ws.max_row, 16).font = Font(color="FFFFFF", bold=True)
    ws2 = wb.create_sheet("Brandkrydstjek")
    ws2.append(["Strækning", "Konstruktion", "Status i dag", "ETA i dag", "Status efter", "ETA efter", "Halv tykkelse mm", "Konflikt", "Anbefaling"])
    for r in audit["runs"]:
        for x in r["fire"]:
            e = lambda v: f"{v['ref']['eta']} {v['ref']['section']} s. {v['ref']['page']}" if v.get("ref") else ""
            ws2.append([f"{r['id']} {r['name']}", f"{x['host']} ({x['rating']})", x["now"]["status"], e(x["now"]), x["new"]["status"], e(x["new"]),
                        x["half_thickness_mm"], "ja" if x["conflict"] else "nej", x["recommendation"]])
    ws3 = wb.create_sheet("Materialer")
    ws3.append(["Produkt", "Rør Ø mm", "Tykkelse mm", "Kappe", "Meter", "Bestilles m", "Stk."])
    for b in audit["bom"]:
        ws3.append([b["product"], b["pipe_od_mm"], b["thickness_mm"], b["cladding"], b.get("metres"), b.get("order_m"), b.get("pieces")])
    for w in (ws, ws2, ws3):
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(out)


if __name__ == "__main__":
    audit = json.loads((ROOT / "results" / "audit.json").read_text(encoding="utf-8"))
    val = json.loads((ROOT / "results" / "validation.json").read_text(encoding="utf-8"))
    (ROOT / "out").mkdir(exist_ok=True)
    pdf(audit, val, ROOT / "out" / "Isoleringsrapport_blok_c.pdf")
    xlsx(audit, ROOT / "out" / "Isoleringsaudit_blok_c.xlsx")
    for key, name, cols in (("V1_DS452_annexA", "validering_ds452_anneks_a.csv", ["table", "D", "class", "published", "model"]),
                            ("V2_Energihaandbogen", "validering_energihaandbogen.csv", ["table", "D", "from_mm", "to_mm", "published", "model", "model_with_halving"])):
        with open(ROOT / "out" / name, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(cols)
            for d in val[key]["data"]:
                w.writerow([d.get(c) for c in cols])
    print("report written")
