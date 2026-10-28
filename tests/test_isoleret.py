"""Isoleret acceptance tests. Run from the product root: python -m pytest tests"""
import functools
import http.server
import json
import os
import math
import random
import sys
import threading
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from isoleret import audit, ds452, physics, survey, validate  # noqa: E402

WEB = ROOT / "web" if (ROOT / "web").exists() else ROOT / "docs" / "app"


# ---------- I1 physics ----------

def test_i1_physics_limits_and_monotonicity():
    # thicker insulation -> lower loss; loss proportional to Δθ for constant λ
    qs = [physics.pipe(0.06, [(t / 1000, "ref_ens")], 55, 15)["q"] for t in (10, 20, 40, 80, 160)]
    assert all(a > b for a, b in zip(qs, qs[1:]))
    # very large diameter approaches the plane-wall solution per m² of surface
    D, t = 5.0, 0.05
    q_cyl = physics.pipe(D, [(t, "ref_ens")], 60, 20)["q"] / (math.pi * (D + 2 * t))
    q_flat = physics.flat([(t, "ref_ens")], 60, 20, orientation="horizontal")["q"]
    assert abs(q_cyl - q_flat) / q_flat < 0.06
    # energy balance: conduction through the layer equals surface loss
    r = physics.pipe(0.042, [(0.03, "stenuld_roerskaal")], 60, 20, "pvc")
    assert abs(r["q"] - (r["ts"] - 20) * r["h_se"] * math.pi * r["De"]) < 1e-3 * r["q"]
    # λ taken at the layer mean temperature, inside the material curve
    assert 0.034 < r["lam"][0] < 0.039


def test_i1b_dew_point_and_condensation():
    assert abs(physics.dew_point(20, 50) - 9.3) < 0.1 and abs(physics.dew_point(22, 55) - 12.5) < 0.2
    t = physics.condensation_thickness(0.054, "elastomer", 10, 22, 55, "elastomer")
    assert t is not None and physics.pipe(0.054, [(t, "elastomer")], 10, 22, "elastomer")["ts"] >= physics.dew_point(22, 55) - 1e-6
    assert physics.condensation_thickness(0.054, "elastomer", 10, 12, 65) == 0.0       # dew point 5.5 °C < 10 °C water


# ---------- I2 DS 452 ----------

def test_i2_table_6_4_and_limits():
    assert ds452.required_class("dhw", "distribution", None, 20)["class"] == 4
    assert ds452.required_class("dhw", "distribution", None, 4)["class"] == 6
    assert ds452.required_class("heating", "distribution", "other", 12)["class"] == 4
    assert ds452.required_class("heating", "distribution", "season_regulated", 12)["class"] == 3
    assert ds452.required_class("heating", "connection_other_room", None, 20)["class"] == 2
    assert ds452.required_class("cold_water", "pipes", None, 4)["class"] == "DS 439"
    assert ds452.required_class("cooling", "coils_0_15", None, 20)["condensation"]
    assert abs(ds452.u_limit(4, 0.06) - 0.25) < 1e-9 and abs(ds452.u_limit(6, 0.018) - 0.1353) < 1e-9
    assert ds452.support_factor("traditional") == 1.15 and ds452.support_factor("traditional", outside=True) == 1.25


# ---------- I3/I4 independent references ----------

@pytest.fixture(scope="module")
def val():
    return validate.run()


def test_i3_ds452_annex_a(val):
    v = val["V1_DS452_annexA"]
    assert v["cells"] == 72 and v["rms_mm"] < 1.0 and v["max_abs_mm"] <= 3.0 and v["within_2mm"] >= 68
    assert abs(v["fitted_lambda"] - 0.0459) < 0.0006                                 # stable design λ (materials.ref_ds452)


def test_i4_energihaandbogen(val):
    v = val["V2_Energihaandbogen"]
    assert v["cells"] == 96 and v["bare_max_rel_pct"] < 3.0 and v["insulated_mean_abs_kwh"] < 5.0
    h = v["halving_finding"]
    assert abs(h["mean_rel_err_without_halving_pct"]) < 3 and h["mean_rel_err_with_halving_pct"] < -40   # the 50 % rule is not in the numbers


# ---------- I5 audit on the demonstration model ----------

@pytest.fixture(scope="module")
def result():
    return audit.run()


def test_i5_answer_key(result):
    ev = result["evaluation"]
    assert ev["runs"] == 25 and ev["all_ok"], [r for r in ev["rows"] if not r["ok"]]


def test_i6_audit_consistency(result):
    s, runs = result["summary"], result["runs"]
    assert s["kwh_saved"] == sum(r["proposal"]["kwh"] for r in runs if r["proposal"])
    for r in runs:
        p = r["proposal"]
        if p:
            assert p["U_new"] <= r["U_limit"] + 1e-9 if r["U_limit"] else True              # every proposal meets the class
            assert p["q_new"] <= r["q_now"] + 1e-9 or r["tm"] < r["ambient"]
        if r["status"] == "KONDENS":
            assert p and p["ts_new"] >= r["dew_point"]
    assert all(b.get("order_m", 1) >= b.get("metres", 0) for b in result["bom"])


def test_i7_fire_crosscheck(result):
    pens = [(r, x) for r in result["runs"] for x in r["fire"]]
    assert len(pens) == 29                                                               # 8 risers x 3 decks + 5 pipes through the plant-room wall
    riser = [x for r, x in pens if r["id"] == "S-A1-V"]
    assert all(x["now"]["status"] == "TESTET" and x["conflict"] and x["keep_tested"] for x in riser)
    assert all(x["now"]["ref"]["section"].startswith("A.4.") and x["now"]["ref"]["insulation_length"] == 1000 for x in riser)
    assert all("tabel 6.2" in x["recommendation"] for x in riser)
    k1 = [x for r, x in pens if r["id"] == "K1"][0]
    assert k1["new"]["status"] == "TESTET" and not k1["conflict"]


def test_i8_survey_round_trip(result, tmp_path):
    from isoleret import takeoff
    runs, _ = takeoff.read(audit.MODEL)
    p = tmp_path / "s.csv"
    survey.write(runs, p)
    back = survey.read(p)
    a = {r["id"]: (r["status"], r["class"], r["proposal"]["thickness_mm"] if r["proposal"] else None) for r in result["runs"]}
    b = {r["id"]: (lambda x: (x["status"], x["class"], x["proposal"]["thickness_mm"] if x["proposal"] else None))(audit.audit_run(r)) for r in back}
    assert a == b


# ---------- I9 browser parity ----------

def _serve():
    h = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(WEB))
    h.log_message = lambda *a, **k: None
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), h)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


def test_i9_browser_parity():
    from playwright.sync_api import sync_playwright
    rng = random.Random(5)
    mats = ["stenuld_roerskaal", "stenuld_lamel", "elastomer", "pe_skum", "pir"]
    cases = []
    for _ in range(400):
        D = rng.choice([0.015, 0.022, 0.028, 0.0424, 0.0603, 0.0889, 0.1143])
        tm, ta = rng.choice([(60, 20), (55, 12), (70, 22), (45, 4), (10, 22), (6, 25), (35, 15)])
        m = rng.choice(mats)
        t = rng.choice([0, 0.009, 0.02, 0.03, 0.05, 0.08, 0.12])
        clad = rng.choice(["laerred", "pvc", "alu_oxid", "elastomer"])
        cases.append({"D": D, "tm": tm, "ta": ta, "m": m, "t": t, "clad": clad, "bare": rng.choice(["steel", "copper"]), "cls": rng.randint(1, 6), "rh": 60})
    py = []
    for c in cases:
        r = physics.pipe(c["D"], [(c["t"], c["m"])] if c["t"] else [], c["tm"], c["ta"], c["clad"], bare=c["bare"])
        need = physics.thickness_for(c["D"], c["m"], c["tm"], c["ta"], ds452.u_limit(c["cls"], c["D"]), c["clad"], 1.15)
        cond = physics.condensation_thickness(c["D"], c["m"], c["tm"], c["ta"], c["rh"], c["clad"]) if c["tm"] < c["ta"] else None
        py.append([r["q"], r["ts"], need, cond])
    kb = json.loads((WEB / "data" / "kb.json").read_text(encoding="utf-8"))
    srv = _serve()
    try:
        with sync_playwright() as pw:
            b = pw.chromium.launch(channel=os.environ.get("PW_CHANNEL", "msedge") or None, headless=True)
            pg = b.new_page()
            pg.set_content('<meta charset="utf-8">')
            pg.add_script_tag(content=(WEB / "engine.js").read_text(encoding="utf-8"))
            pg.evaluate("kb => IEngine.load(kb)", kb)
            js = pg.evaluate("""cs => cs.map(c => { const E = IEngine; const r = E.pipe(c.D, c.t ? [[c.t, c.m]] : [], c.tm, c.ta, c.clad, {bare: c.bare});
                const need = E.thicknessFor(c.D, c.m, c.tm, c.ta, E.uLimit(c.cls, c.D), c.clad, 1.15);
                const cond = c.tm < c.ta ? E.condensationThickness(c.D, c.m, c.tm, c.ta, c.rh, c.clad) : null; return [r.q, r.ts, need, cond]; })""", cases)
            b.close()
    finally:
        srv.shutdown()
    bad = []
    for i, (a, bb) in enumerate(zip(py, js)):
        for x, y in zip(a, bb):
            if (x is None) != (y is None) or (x is not None and abs(x - y) > 1e-9 * max(1, abs(x))):
                bad.append((i, cases[i], a, bb))
                break
    (ROOT / "results" / "parity.json").write_text(json.dumps({"cases": len(cases), "identical_1e-9": len(cases) - len(bad), "first": bad[:2]}, indent=1), encoding="utf-8")
    assert not bad, bad[:1]


# ---------- I10 deliverables ----------

def test_i10_deliverables():
    out = ROOT / "out" if (ROOT / "out").exists() else WEB / "downloads"
    import pymupdf
    from openpyxl import load_workbook
    d = pymupdf.open(out / "Isoleringsrapport_blok_c.pdf")
    assert d.page_count >= 5 and "Brandkrydstjek" in "".join(p.get_text() for p in d)
    wb = load_workbook(out / "Isoleringsaudit_blok_c.xlsx")
    assert wb["Strækninger"].max_row - 1 == 25 and wb["Brandkrydstjek"].max_row - 1 == 29


def test_i11_calculation_trace_matches_audit():
    """'Vis beregning' recomputes every run step by step in the browser; each must land on the audit's number (✓, never ✗)."""
    from playwright.sync_api import sync_playwright
    runs = json.loads((WEB / "data" / "audit.json").read_text(encoding="utf-8"))["runs"]
    srv = _serve()
    try:
        with sync_playwright() as pw:
            b = pw.chromium.launch(channel=os.environ.get("PW_CHANNEL", "msedge") or None, headless=True, args=["--use-angle=swiftshader", "--enable-unsafe-swiftshader"])
            pg = b.new_page()
            pg.goto(f"http://127.0.0.1:{srv.server_address[1]}/index.html")
            pg.wait_for_function("window.IFCalc && window.IEngine && IEngine.kb()")
            res = pg.evaluate("rs => rs.map(r => { const h = IFCalc.trace(r); return [r.id, (h.match(/okline\">✓/g) || []).length, (h.match(/badline\">✗/g) || []).length]; })", runs)
            b.close()
    finally:
        srv.shutdown()
    for rid, ok, bad in res:
        assert bad == 0 and ok >= 1, (rid, ok, bad)


def test_i12_reproduce():
    from isoleret import reproduce
    o = reproduce.run()
    assert o["ok"], [c for c in o["checks"] if not c["ok"]]
